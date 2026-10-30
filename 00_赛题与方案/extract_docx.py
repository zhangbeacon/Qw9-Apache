import zipfile
from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

path = r"E:\强网拟态\qw9-design-challenge-zh.docx"
out = r"E:\强网拟态\docx_content.txt"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

lines = []


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


doc = Document(path)

lines.append("=== CORE PROPERTIES ===")
cp = doc.core_properties
for attr in ["title", "subject", "author", "last_modified_by", "created", "modified", "revision", "comments"]:
    try:
        lines.append(f"{attr}: {getattr(cp, attr)}")
    except Exception as e:
        lines.append(f"{attr}: <error {e}>")

lines.append("\n=== BODY ===")
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        text = block.text
        if text.strip():
            lines.append(f"[{block.style.name}] {text}")
    else:
        lines.append("[TABLE]")
        for row in block.rows:
            cells = [c.text.replace("\n", " / ") for c in row.cells]
            lines.append(" | ".join(cells))
        lines.append("[/TABLE]")

lines.append("\n=== HEADERS / FOOTERS ===")
for i, section in enumerate(doc.sections, 1):
    for name in ["header", "footer"]:
        hf = getattr(section, name)
        for p in hf.paragraphs:
            if p.text.strip():
                lines.append(f"[section {i} {name}] {p.text}")
        for t in hf.tables:
            for row in t.rows:
                lines.append(f"[section {i} {name} table] " + " | ".join(c.text for c in row.cells))

lines.append("\n=== ALL w:t TEXT (document.xml, includes text boxes) ===")
with zipfile.ZipFile(path) as z:
    names = z.namelist()
    root = etree.fromstring(z.read("word/document.xml"))
    texts = root.xpath("//w:t/text()", namespaces={"w": W[1:-1]})
    lines.append("".join(texts))

    for part in ["word/comments.xml", "word/header1.xml", "word/footer1.xml"]:
        if part in names:
            lines.append(f"\n=== {part} ===")
            r = etree.fromstring(z.read(part))
            lines.append("".join(r.xpath("//w:t/text()", namespaces={"w": W[1:-1]})))

    lines.append("\n=== PACKAGE PARTS ===")
    for n in names:
        if n.startswith("word/") and n.endswith(".xml"):
            lines.append(n)

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print("written:", out)
