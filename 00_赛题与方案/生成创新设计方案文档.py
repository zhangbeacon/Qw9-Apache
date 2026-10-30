from pathlib import Path
import json
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(r"E:\强网拟态\00_赛题与方案")
SOURCE = ROOT / "第一阶段_创新设计理念方案_初稿.md"
OUTPUT = ROOT / "第一阶段_创新设计理念方案_初稿.docx"


def set_font(font, size=None, bold=None, color=None, name="Microsoft YaHei"):
    font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)


def set_east_asia_style(style, font_name="Microsoft YaHei"):
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), font_name)


def set_east_asia_run(run, font_name="Microsoft YaHei"):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), font_name)


def add_hyperlink(paragraph, text, target):
    part = paragraph.part
    rel_id = part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245A82")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), "Microsoft YaHei")
    rpr.append(fonts)
    run.append(rpr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")


def add_inline(paragraph, text, size=None, color=None):
    pos = 0
    for match in TOKEN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_east_asia_run(run)
            if size is not None:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        link = LINK.match(token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        else:
            value = token[2:-2] if token.startswith("**") else token[1:-1]
            run = paragraph.add_run(value)
            set_east_asia_run(run)
            if token.startswith("**"):
                run.bold = True
            elif token.startswith("`"):
                run.font.name = "Consolas"
            if size is not None:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_east_asia_run(run)
        if size is not None:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    if pos == 0 and not text:
        paragraph.add_run("")


def shade(element, color):
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:fill"), color)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    tr_pr.append(flag)


def set_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def add_table(doc, raw_rows):
    rows = []
    for raw in raw_rows:
        cells = [v.strip() for v in raw.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    count = max(map(len, rows))
    rows = [r + [""] * (count - len(r)) for r in rows]
    table = doc.add_table(rows=0, cols=count)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if count == 2:
        widths = [3.5, 13.1]
    elif count == 3:
        widths = [3.0, 6.3, 7.3]
    elif count == 4:
        widths = [4.0, 4.2, 4.2, 4.2]
    elif count == 5:
        widths = [3.1, 2.3, 2.3, 2.3, 6.6]
    else:
        widths = [16.6 / count] * count
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for idx, row_values in enumerate(rows):
        row = table.add_row()
        cant_split(row)
        if idx == 0:
            set_table_header(row)
        for i, (cell, value) in enumerate(zip(row.cells, row_values)):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell._tc, "284B63" if idx == 0 else ("F4F7FA" if idx % 2 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, value, size=9, color="FFFFFF" if idx == 0 else None)
            if idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_line(doc, line):
    p = doc.add_paragraph(style="Code")
    shade(p._p, "F0F3F6")
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(line if line else " ")
    set_east_asia_run(r)
    r.font.size = Pt(8.4)
    r.font.name = "Consolas"


def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal.font, size=10.3, color="22313D")
    set_east_asia_style(normal)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    for label, size, before, after in [("Title", 20, 0, 10), ("Subtitle", 12, 0, 10), ("Heading 1", 14, 12, 5), ("Heading 2", 11.5, 9, 4), ("Heading 3", 10.5, 6, 3)]:
        style = doc.styles[label]
        set_east_asia_style(style)
        set_font(style.font, size=size, bold=label.startswith("Heading") or label == "Title", color="173B55" if label != "Subtitle" else "567082")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if label.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    for label in ("List Bullet", "List Number"):
        style = doc.styles[label]
        set_east_asia_style(style)
        set_font(style.font, size=10.3, color="22313D")
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2
    if "Code" not in doc.styles:
        from docx.enum.style import WD_STYLE_TYPE
        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code"]
    set_east_asia_style(code)
    set_font(code.font, size=8.4, color="273746")


def setup_sections(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("赛题一 · 第一阶段设计理念（讨论稿）")
    set_east_asia_run(r)
    set_font(r.font, size=8, color="778899")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("讨论初稿 · 第 ")
    set_east_asia_run(r)
    set_font(r.font, size=8, color="778899")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    r = footer.add_run(" 页")
    set_east_asia_run(r)
    set_font(r.font, size=8, color="778899")


def main():
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    setup_styles(doc)
    setup_sections(doc)
    doc.core_properties.title = "面向攻击链阻断的互补式异构变体生成与验证系统"
    doc.core_properties.subject = "第一阶段设计理念方案讨论稿"
    lines = text.splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_code_line(doc, line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line[2:])
        elif line.startswith("## 第一阶段"):
            p = doc.add_paragraph(style="Subtitle")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line[3:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:])
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade(p._p, "EAF2F7")
            add_inline(p, line[2:])
        elif line.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "□ " + line[6:])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            numbered = re.match(r"^\d+\.\s+(.+)$", line)
            if numbered:
                p = doc.add_paragraph(style="List Number")
                add_inline(p, numbered.group(1))
            else:
                p = doc.add_paragraph()
                add_inline(p, line)
        i += 1
    if in_code:
        raise ValueError("Unclosed fenced code block")
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    all_text = "\n".join(p.text for p in check.paragraphs) + "\n" + "\n".join(c.text for t in check.tables for row in t.rows for c in row.cells)
    required = ["互补式异构变体", "待选择", "尚未完成", "至少三个", "共同攻击成功比例", "攻击链", "证据卡", "伪代码", "功能一致性", "团队分工", "待补充"]
    missing = [x for x in required if x not in all_text]
    if missing:
        raise AssertionError(f"Required content missing: {missing}")
    print(json.dumps({"output": str(OUTPUT), "paragraphs": len(check.paragraphs), "tables": len(check.tables), "bytes": OUTPUT.stat().st_size, "required_checks": len(required)}, ensure_ascii=True))


if __name__ == "__main__":
    main()
